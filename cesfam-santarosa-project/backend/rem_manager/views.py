# rem_manager/views.py
import json
import re

import openpyxl
from rest_framework import serializers, status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from authentication.models import RegistroAuditoria
from authentication.permissions import EsGestorOSubdireccion
from .models import RemArchivo

# Meses tal como aparecen en la hoja "NOMBRE" de la Serie A (MINSAL - DEIS)
MESES = {
    'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5, 'JUNIO': 6,
    'JULIO': 7, 'AGOSTO': 8, 'SEPTIEMBRE': 9, 'OCTUBRE': 10,
    'NOVIEMBRE': 11, 'DICIEMBRE': 12,
}

# Hojas REM válidas: A01, A02, ..., A34 (pueden tener sufijos como A11a, A30AR)
HOJA_REM_RE = re.compile(r'^A\d{2}', re.IGNORECASE)

EXTENSIONES_PERMITIDAS = ('.xlsx', '.xlsm')


class RemUploadView(APIView):
    """
    RF-03: Carga de información REM.
    Valida que la planilla sea una Serie A oficial del MINSAL-DEIS y que
    el período declarado coincida con el período del archivo. Registra la
    operación en la bitácora de auditoría (RF-12).
    """
    parser_classes = [MultiPartParser, FormParser]
    permission_classes = [IsAuthenticated, EsGestorOSubdireccion]

    def post(self, request):
        archivo = request.FILES.get('archivo')
        if archivo is None:
            return Response(
                {'error': 'Debe adjuntar un archivo Excel.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        nombre_archivo = archivo.name.lower()
        if not nombre_archivo.endswith(EXTENSIONES_PERMITIDAS):
            self._registrar_rechazo(request, archivo.name,
                                     'Formato de archivo no permitido (solo .xlsx o .xlsm).')
            return Response(
                {'error': 'Formato no permitido. Suba la planilla en formato .xlsx o .xlsm.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            periodo_mes = int(request.data.get('periodo_mes'))
            periodo_year = int(request.data.get('periodo_year'))
        except (TypeError, ValueError):
            return Response(
                {'error': 'Debe indicar el mes y el año del período.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # --- Validación de la estructura de la planilla REM Serie A ---
        try:
            libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        except Exception:
            self._registrar_rechazo(request, archivo.name,
                                     'El archivo no es un Excel válido o está dañado.',
                                     periodo_mes, periodo_year)
            return Response(
                {'error': 'El archivo no se pudo abrir. Verifique que sea una planilla REM válida.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if 'NOMBRE' not in libro.sheetnames:
            self._registrar_rechazo(request, archivo.name,
                                     'La planilla no contiene la hoja "NOMBRE": no es una Serie A oficial.',
                                     periodo_mes, periodo_year)
            return Response(
                {'error': 'La planilla no corresponde a un REM Serie A oficial (falta hoja "NOMBRE").'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        hoja_nombre = libro['NOMBRE']
        # Las etiquetas pueden cambiar de fila entre versiones de la planilla:
        # se buscan en la hoja en lugar de usar celdas fijas.
        mes_archivo = ''
        year_archivo_raw = None
        establecimiento = ''
        for fila in hoja_nombre.iter_rows(min_row=1, max_row=20, max_col=6):
            etiqueta = str(fila[0].value or '').strip().upper()
            if etiqueta == 'MES:' and mes_archivo == '':
                mes_archivo = str(fila[1].value or '').strip().upper()
            elif etiqueta == 'AÑO:' and year_archivo_raw is None:
                year_archivo_raw = fila[1].value
            elif etiqueta.startswith('NOMBRE ESTABLECIMIENTO'):
                establecimiento = str(fila[1].value or '').strip()

        try:
            year_archivo = int(float(str(year_archivo_raw).strip()))
        except (TypeError, ValueError):
            self._registrar_rechazo(request, archivo.name,
                                     'No se pudo leer el año de la planilla (etiqueta "AÑO:" de la hoja NOMBRE).',
                                     periodo_mes, periodo_year)
            return Response(
                {'error': 'No se pudo leer el año de la planilla (etiqueta "AÑO:" de la hoja NOMBRE).'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if mes_archivo not in MESES:
            self._registrar_rechazo(request, archivo.name,
                                     f'Mes ilegible en la planilla: "{mes_archivo}".',
                                     periodo_mes, periodo_year)
            return Response(
                {'error': f'No se pudo leer el mes de la planilla (etiqueta "MES:" de la hoja NOMBRE): "{mes_archivo}".'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        mes_archivo_num = MESES[mes_archivo]

        # El período declarado por el usuario debe coincidir con el del archivo
        if mes_archivo_num != periodo_mes or year_archivo != periodo_year:
            detalle = (f'El archivo corresponde a {mes_archivo.capitalize()} {year_archivo}, '
                       f'pero se seleccionó {periodo_mes:02d}-{periodo_year}.')
            self._registrar_rechazo(request, archivo.name, detalle,
                                    periodo_mes, periodo_year)
            return Response({'error': detalle}, status=status.HTTP_400_BAD_REQUEST)

        hojas_rem = [s for s in libro.sheetnames if HOJA_REM_RE.match(s)]
        if not hojas_rem:
            detalle = 'La planilla no contiene hojas REM (A01, A02, ...).'
            self._registrar_rechazo(request, archivo.name, detalle,
                                    periodo_mes, periodo_year)
            return Response({'error': detalle}, status=status.HTTP_400_BAD_REQUEST)

        # Recarga del período: el archivo anterior queda marcado como reemplazado
        rempleazados = RemArchivo.objects.filter(
            periodo_mes=periodo_mes,
            periodo_year=periodo_year,
            estado=RemArchivo.Estado.VALIDO,
        ).update(estado=RemArchivo.Estado.REEMPLAZADO)

        # openpyxl dejó el puntero de lectura al final: rebobinar antes de guardar
        archivo.seek(0)
        registro = RemArchivo.objects.create(
            usuario=request.user,
            archivo=archivo,
            archivo_nombre=archivo.name,
            periodo_mes=periodo_mes,
            periodo_year=periodo_year,
            establecimiento=establecimiento,
            hojas_rem=len(hojas_rem),
            estado=RemArchivo.Estado.VALIDO,
            observaciones=(
                f'Planilla validada: {len(hojas_rem)} hojas REM detectadas'
                + (' (recarga: reemplaza carga anterior)' if rempleazados else '')
            ),
        )

        # RF-12: auditoría de la carga
        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion='CARGA_REM',
            tabla_afectada='RemArchivo',
            valor_nuevo=json.dumps({
                'archivo': archivo.name,
                'periodo': f'{periodo_mes:02d}-{periodo_year}',
                'establecimiento': establecimiento,
                'hojas_rem': len(hojas_rem),
            }, ensure_ascii=False),
        )

        return Response({
            'mensaje': 'Planilla REM cargada y validada exitosamente.',
            'periodo': f'{periodo_mes:02d}-{periodo_year}',
            'establecimiento': establecimiento,
            'hojas_rem': len(hojas_rem),
            'recarga': bool(rempleazados),
            'id': registro.id,
        }, status=status.HTTP_201_CREATED)

    def _registrar_rechazo(self, request, nombre_archivo, motivo,
                           periodo_mes=None, periodo_year=None):
        """Registra en auditoría y en el historial los intentos de carga rechazados."""
        RemArchivo.objects.create(
            usuario=request.user,
            archivo=None,
            archivo_nombre=nombre_archivo,
            periodo_mes=periodo_mes or 0,
            periodo_year=periodo_year or 0,
            estado=RemArchivo.Estado.RECHAZADO,
            observaciones=motivo,
        )
        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion='CARGA_REM_RECHAZADA',
            tabla_afectada='RemArchivo',
            valor_nuevo=json.dumps({
                'archivo': nombre_archivo,
                'motivo': motivo,
            }, ensure_ascii=False),
        )


class RemArchivoSerializer(serializers.ModelSerializer):
    usuario = serializers.SerializerMethodField()

    class Meta:
        model = RemArchivo
        fields = [
            'id', 'periodo_mes', 'periodo_year', 'archivo_nombre',
            'establecimiento', 'hojas_rem', 'estado', 'observaciones',
            'fecha_carga', 'usuario',
        ]

    def get_usuario(self, obj):
        if obj.usuario:
            return str(obj.usuario)
        return 'Sistema'


class RemArchivoListView(APIView):
    """Historial de cargas REM (válidas, reemplazadas y rechazadas)."""
    permission_classes = [IsAuthenticated, EsGestorOSubdireccion]

    def get(self, request):
        registros = RemArchivo.objects.select_related('usuario')[:200]
        data = RemArchivoSerializer(registros, many=True).data
        return Response({'resultados': data})
